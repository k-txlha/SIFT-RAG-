"""Unit tests for file parsers (TXT, PDF, DOCX)."""

from io import BytesIO
import unittest

from parsers import (
    EmptyDocumentError,
    UnsupportedFileError,
    parse_document,
    parse_docx,
    parse_txt,
)


class TestParsers(unittest.TestCase):
    def test_parse_valid_txt(self):
        content = b"Apache Kafka is a distributed event streaming platform.\n\nIt is fast."
        ext, sections = parse_document("sample.txt", content)
        self.assertEqual(ext, "txt")
        self.assertEqual(len(sections), 1)
        self.assertIn("Apache Kafka", sections[0].text)
        self.assertIsNone(sections[0].page_number)

    def test_parse_empty_txt(self):
        content = b"   \n\n\t  "
        with self.assertRaises(EmptyDocumentError):
            parse_document("empty.txt", content)

    def test_unsupported_file_extension(self):
        with self.assertRaises(UnsupportedFileError):
            parse_document("image.png", b"fake image content")

    def test_parse_docx(self):
        import docx
        doc = docx.Document()
        doc.add_heading("System Architecture", level=1)
        doc.add_paragraph("This is a test paragraph for DOCX parsing.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"

        buffer = BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()

        ext, sections = parse_document("arch.docx", content)
        self.assertEqual(ext, "docx")
        self.assertTrue(len(sections) >= 1)
        self.assertIn("System Architecture", sections[0].text)
        self.assertIn("test paragraph", sections[0].text)
        self.assertIn("Header 1", sections[0].text)

    def test_parse_empty_docx(self):
        import docx
        doc = docx.Document()
        buffer = BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()

        with self.assertRaises(EmptyDocumentError):
            parse_docx(content)


if __name__ == "__main__":
    unittest.main()
